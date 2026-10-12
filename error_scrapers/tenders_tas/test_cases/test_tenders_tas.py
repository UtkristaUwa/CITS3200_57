"""
Tests for the Tasmanian eTendering scraper, same discipline as
error_scrapers/tenders_wa/test_cases/test_tenders_wa.py.

All six fixtures are live captures:
  tenders_tas_list.html           Open For Bids list (every open tender on one page)
  tenders_tas_details.html        one tender's detail page (6652)
  tenders_tas_attachements.html   the Downloadables page of a tender with four files
  tenders_tas_login.html          the login page
  tenders_tas_fail_login.html     the login page after a wrong password
  tenders_tas_success_login.html  the home page after a successful login
The files themselves are built in-test, and the fake portal answers each step
the way the real one did when the pages were captured.
"""

import io
import os
from urllib.parse import parse_qs

import docx
import fitz
import httpx
import pytest

from error_scrapers import common, run_notes
from error_scrapers.tenders_tas import scraper

FIXTURES = os.path.dirname(__file__)

BASE = "https://www.tenders.tas.gov.au"
REFERENCE = "NRET MR 2026-08"
DETAIL_ID = "6652"
DOCS_ID = "6649"  # the tender the Downloadables fixture was captured from
TOKEN = "SCnzHPizet5SMTSKWkd9"


def _read_fixture(name: str) -> str:
    with open(os.path.join(FIXTURES, name), "r", encoding="utf-8") as f:
        return f.read()


def _listing() -> str:
    return _read_fixture("tenders_tas_list.html")


def _detail() -> str:
    return _read_fixture("tenders_tas_details.html")


def _documents_page() -> str:
    return _read_fixture("tenders_tas_attachements.html")


def _pdf_bytes(text: str) -> bytes:
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def _xlsx_bytes() -> bytes:
    import openpyxl
    buf = io.BytesIO()
    book = openpyxl.Workbook()
    book.active["A1"] = "rates"
    book.save(buf)
    return buf.getvalue()


def _docx_bytes(text: str) -> bytes:
    buf = io.BytesIO()
    document = docx.Document()
    document.add_paragraph(text)
    document.save(buf)
    return buf.getvalue()


# What each /Tender/DownloadFile/<id> in the Downloadables fixture returns.
FILES = {
    "119017": ("application/pdf", lambda: _pdf_bytes("Volume one text")),
    "119018": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",
               lambda: _docx_bytes("Tender form text")),
    "119019": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", lambda: b"not a real xlsx"),
    "119020": ("application/pdf", lambda: _pdf_bytes("Volume two text")),
}


def _portal(list_html=None, detail_html=None, documents_html=None, password="secret",
            files=None, lapse_after_logins=None):
    """
    A fake Tasmanian portal. Serves the list, any tender's detail page (the Unique
    Tender ID rewritten per site id so tenders do not share a folder), the
    Downloadables page and the files. Documents and files need a prior successful
    login; `lapse_after_logins` makes the session lapse (sign-in page served) until
    the client has logged in that many times.
    """
    state = {"requests": [], "logins": 0, "failed_logins": 0, "signed_in": False, "posts": []}
    files = FILES if files is None else files

    def signed_out_page():
        return httpx.Response(200, text=_read_fixture("tenders_tas_login.html"))

    def handler(request: httpx.Request) -> httpx.Response:
        path = request.url.path
        state["requests"].append((request.method, path))
        if path == "/Account/Logon":
            if request.method == "GET":
                return httpx.Response(200, text=_read_fixture("tenders_tas_login.html"))
            form = parse_qs(request.content.decode())
            state["posts"].append(form)
            if form.get("UserName") == ["user"] and form.get("Password") == [password]:
                state["logins"] += 1
                state["signed_in"] = True
                return httpx.Response(200, text=_read_fixture("tenders_tas_success_login.html"))
            state["failed_logins"] += 1
            return httpx.Response(200, text=_read_fixture("tenders_tas_fail_login.html"))
        if path == "/OpenForBids/List/Public/ClosingDate":
            return httpx.Response(200, text=list_html if list_html is not None else _listing())
        if path.startswith("/OpenForBids/Details/"):
            site_id = path.rsplit("/", 1)[1]
            html = detail_html if detail_html is not None else _detail()
            return httpx.Response(200, text=html.replace(REFERENCE, f"TAS{site_id}"))
        lapsed = lapse_after_logins is not None and state["logins"] < lapse_after_logins
        if path.startswith("/OpenForBids/Downloadables/"):
            if not state["signed_in"] or lapsed:
                return signed_out_page()
            return httpx.Response(200, text=documents_html if documents_html is not None else _documents_page())
        if path.startswith("/Tender/DownloadFile/"):
            if not state["signed_in"] or lapsed:
                return signed_out_page()
            kind, make = files[path.rsplit("/", 1)[1]]
            return httpx.Response(200, content=make(), headers={"content-type": kind})
        return httpx.Response(404)

    transport = httpx.MockTransport(handler)
    transport.state = state
    return transport


def _client(transport):
    return httpx.Client(transport=transport, follow_redirects=True)


def _patch_client(monkeypatch, transport):
    real_client = httpx.Client
    monkeypatch.setattr(
        scraper.httpx, "Client",
        lambda **kwargs: real_client(transport=transport, **kwargs),
    )


def _set_credentials(monkeypatch, password="secret"):
    monkeypatch.setenv("TAS_USERNAME", "user")
    monkeypatch.setenv("TAS_PASSWORD", password)


@pytest.fixture(autouse=True)
def _no_pause(monkeypatch):
    monkeypatch.setattr(scraper, "PAUSE_SECONDS", 0)
    monkeypatch.setattr(common, "PAGE_RETRY_SECONDS", 0)
    monkeypatch.setattr(common, "DOWNLOAD_RETRY_SECONDS", 0)
    monkeypatch.delenv("TAS_USERNAME", raising=False)
    monkeypatch.delenv("TAS_PASSWORD", raising=False)
    run_notes.reset()


# ---------------------------------------------------------------------------
# Listing
# ---------------------------------------------------------------------------

def test_listing_finds_every_real_tender_once():
    entries = scraper.parse_listing(_listing())
    ids = [e["id"] for e in entries]
    assert len(entries) == 39
    assert len(set(ids)) == len(ids)
    assert entries[0]["id"] == DETAIL_ID
    assert entries[0]["url"] == f"{BASE}/OpenForBids/Details/{DETAIL_ID}"


def test_listing_skips_the_demonstration_notice():
    titles = [e["title"] for e in scraper.parse_listing(_listing())]
    assert not any(t.lower().startswith("demonstration notice") for t in titles)


def test_listing_row_has_its_title_agency_category_and_dates():
    first = scraper.parse_listing(_listing())[0]
    assert first["title"].startswith("East Coast Rock Lobster Translocation Program Services Panel")
    assert first["agency"] == "Department of Natural Resources and Environment Tasmania"
    assert first["category"] == "Farming and Fishing and Forestry and Wildlife Contracting Services"
    assert first["opened"] == "26/09/2026"
    assert first["closes"] == "12/10/2026, at 05:00 PM"


def test_a_closing_heading_applies_to_every_tender_under_it():
    entries = {e["id"]: e for e in scraper.parse_listing(_listing())}
    # Two tenders sit under the 14/10/2026 heading in the capture.
    assert entries["6649"]["closes"] == "14/10/2026, at 02:00 PM"
    assert sum(1 for e in entries.values() if e["closes"] == "14/10/2026, at 02:00 PM") == 2


def test_perpetually_open_tenders_are_kept():
    entries = {e["id"]: e for e in scraper.parse_listing(_listing())}
    assert entries["5128"]["closes"] == "Perpetually Open"


def test_listing_with_no_rows_is_empty():
    assert scraper.parse_listing("<html><body>nothing</body></html>") == []


# ---------------------------------------------------------------------------
# Detail page
# ---------------------------------------------------------------------------

def test_detail_fields():
    fields, code = scraper.parse_detail(_detail())
    assert code == common.SITE_SUCCESS
    assert fields["reference"] == REFERENCE
    assert fields["title"] == "East Coast Rock Lobster Translocation Program Services Panel (2026–27)"
    labelled = fields["labelled"]
    assert labelled["Status"] == "Current"
    assert labelled["Type/Procurement Method"] == "Request for Quotation"
    assert labelled["Closing Date and Time"] == "Monday 12/10/2026, at: 05:00 PM"
    assert labelled["Agency"].splitlines()[-1] == "HOBART TAS 7000"
    assert "translocation" in labelled["Description"].lower()
    assert labelled["Conditions for Participation"].startswith("Holding a current Tasmanian")


def test_a_label_repeated_with_the_same_text_is_kept_once():
    labelled = scraper.parse_detail(_detail())[0]["labelled"]
    assert "---" not in labelled["Location"]
    assert labelled["Location"] == "State-wide"


def test_contact_email_keeps_its_words_on_one_line():
    labelled = scraper.parse_detail(_detail())[0]["labelled"]
    assert "Email: Sonja.Hempel@nre.tas.gov.au" in labelled["Enquiries"]


def test_the_documents_are_available_from_block_does_not_repeat_the_enquiries():
    labelled = scraper.parse_detail(_detail())[0]["labelled"]
    assert "Documents are available from" not in labelled


def test_detail_text_leaves_out_link_addresses_and_button_text():
    fields, _ = scraper.parse_detail(_detail())
    text = scraper.format_detail_text(fields, {})
    assert "/OpenForBids/" not in text and "Mailto" not in text and "http" not in text
    assert "Lodge" not in text and "Download Files" not in text


def test_detail_page_without_the_expected_fields_is_a_structure_change():
    html = _detail().replace('class="itemrow', 'class="row-CHANGED')
    assert scraper.parse_detail(html) == ({}, common.SITE_STRUCTURE_CHANGE)


def test_detail_text_has_the_listing_dates_and_every_field():
    entry = scraper.parse_listing(_listing())[0]
    fields, _ = scraper.parse_detail(_detail())
    text = scraper.format_detail_text(fields, entry)
    assert "TASMANIA ETENDERING: East Coast Rock Lobster" in text
    assert "Closes: 12/10/2026, at 05:00 PM" in text
    assert "Unique Tender ID: NRET MR 2026-08" in text
    assert "Period of Contract (Estimated): From 30/09/2026 to 31/05/2027" in text
    assert "DESCRIPTION:" in text and "ENQUIRIES:" in text


# ---------------------------------------------------------------------------
# Documents page
# ---------------------------------------------------------------------------

def test_documents_page_lists_every_file_with_its_size():
    documents = scraper.parse_documents(_documents_page(), f"{BASE}/OpenForBids/Downloadables/{DOCS_ID}")
    assert [d["name"] for d in documents] == [
        "RFT 4317 - Statewide Bridge Joint Rehabilitation Works 2026-27 - Volume 1.pdf",
        "RFT 4317 - Volume 1 - Attachment 3 - Tender Form.docx",
        "RFT 4317 - Volume 1 - Attachment 10 - Schedule of Rates.xlsx",
        "RFT 4317 - Statewide Bridge Joint Rehabilitation Works 2026-27 - Volume 2.pdf",
    ]
    assert documents[0]["url"] == f"{BASE}/Tender/DownloadFile/119017"
    assert documents[0]["size"] == "1.04 MB"


def test_documents_page_with_no_files_is_empty():
    assert scraper.parse_documents("<html><body>No documents</body></html>") == []


# ---------------------------------------------------------------------------
# Login
# ---------------------------------------------------------------------------

def test_login_page_is_recognised_by_its_password_box():
    assert scraper.find_login_form(_read_fixture("tenders_tas_login.html")) is not None
    assert scraper.find_login_form(_detail()) is None


def test_signed_in_and_failed_pages_are_told_apart():
    ok = _read_fixture("tenders_tas_success_login.html")
    bad = _read_fixture("tenders_tas_fail_login.html")
    assert scraper.is_signed_in(ok) and not scraper.login_failed(ok)
    assert scraper.login_failed(bad) and not scraper.is_signed_in(bad)
    assert not scraper.is_signed_in(_read_fixture("tenders_tas_login.html"))


def test_login_posts_the_token_user_name_and_password(monkeypatch):
    _set_credentials(monkeypatch)
    transport = _portal()
    assert scraper.login(_client(transport)) is True
    form = transport.state["posts"][0]
    assert form["UserName"] == ["user"] and form["Password"] == ["secret"]
    assert "__RequestVerificationToken" in form
    assert "ForgotEmailAddress" not in form  # the forgotten-password boxes are not sent


def test_login_with_a_wrong_password_is_rejected_once_and_not_retried(monkeypatch):
    _set_credentials(monkeypatch, password="wrong")
    transport = _portal()
    assert scraper.login(_client(transport)) is False
    assert transport.state["failed_logins"] == 1


def test_login_without_credentials_does_not_touch_the_site():
    transport = _portal()
    assert scraper.login(_client(transport)) is False
    assert transport.state["requests"] == []


def test_login_page_without_a_form_is_a_structure_change(monkeypatch):
    _set_credentials(monkeypatch)
    transport = httpx.MockTransport(lambda r: httpx.Response(200, text="<html><body>hi</body></html>"))
    with pytest.raises(common.StructureChangedError):
        scraper.login(_client(transport))


def test_login_response_that_is_neither_success_nor_failure_is_a_structure_change(monkeypatch):
    _set_credentials(monkeypatch)

    def handler(request):
        if request.method == "GET":
            return httpx.Response(200, text=_read_fixture("tenders_tas_login.html"))
        return httpx.Response(200, text="<html><body>something else</body></html>")
    with pytest.raises(common.StructureChangedError):
        scraper.login(_client(httpx.MockTransport(handler)))


def test_credentials_come_from_the_environment(monkeypatch):
    assert scraper.credentials() == (None, None)
    _set_credentials(monkeypatch)
    assert scraper.credentials() == ("user", "secret")


# ---------------------------------------------------------------------------
# Downloading documents
# ---------------------------------------------------------------------------

def _signed_in_client(monkeypatch, **portal_args):
    _set_credentials(monkeypatch)
    transport = _portal(**portal_args)
    client = _client(transport)
    assert scraper.login(client)
    return client, transport


def test_each_file_is_saved_on_its_own_with_its_text(tmp_path, monkeypatch):
    client, _ = _signed_in_client(monkeypatch)
    attachments, advertised, failed = scraper.download_documents(client, DOCS_ID, str(tmp_path))
    assert advertised == 4 and len(attachments) == 4
    names = sorted(os.listdir(tmp_path))
    assert "RFT 4317 - Volume 1 - Attachment 3 - Tender Form.docx" in names
    assert "RFT 4317 - Volume 1 - Attachment 3 - Tender Form.docx.txt" in names
    assert "Volume one text" in (tmp_path / "RFT 4317 - Statewide Bridge Joint Rehabilitation Works 2026-27 - Volume 1.pdf.txt").read_text()
    assert not any(n.endswith(".zip") for n in names)
    assert attachments[0]["file_name"].endswith("Volume 1.pdf") and attachments[0]["size_bytes"] > 0


def test_a_file_that_cannot_be_read_is_counted_failed_but_kept(tmp_path, monkeypatch):
    client, _ = _signed_in_client(monkeypatch)
    _, advertised, failed = scraper.download_documents(client, DOCS_ID, str(tmp_path))
    assert failed == 1  # the fake .xlsx is not a real workbook
    assert (tmp_path / "RFT 4317 - Volume 1 - Attachment 10 - Schedule of Rates.xlsx").exists()


def test_one_file_failing_with_a_404_does_not_stop_the_others(tmp_path, monkeypatch):
    base = _portal()
    _set_credentials(monkeypatch)

    def handler(request):
        if request.url.path == "/Tender/DownloadFile/119018":
            return httpx.Response(404)
        return base.handler(request)

    client = _client(httpx.MockTransport(handler))
    assert scraper.login(client)
    attachments, advertised, failed = scraper.download_documents(client, DOCS_ID, str(tmp_path))
    assert advertised == 4 and len(attachments) == 3 and failed >= 1


def test_a_sign_in_page_in_place_of_the_documents_page_is_a_lapsed_session(tmp_path, monkeypatch):
    _set_credentials(monkeypatch)
    client = _client(_portal())  # never logged in
    with pytest.raises(scraper.SessionLapsedError):
        scraper.download_documents(client, DOCS_ID, str(tmp_path))


def test_a_web_page_in_place_of_a_file_is_a_lapsed_session(tmp_path, monkeypatch):
    base = _portal()
    _set_credentials(monkeypatch)

    def handler(request):
        if request.url.path.startswith("/Tender/DownloadFile/"):
            return httpx.Response(200, text="<html>sign in</html>", headers={"content-type": "text/html"})
        return base.handler(request)

    client = _client(httpx.MockTransport(handler))
    assert scraper.login(client)
    with pytest.raises(scraper.SessionLapsedError):
        scraper.download_documents(client, DOCS_ID, str(tmp_path))


def test_a_file_that_fails_with_a_503_is_tried_again(tmp_path, monkeypatch):
    base = _portal()
    _set_credentials(monkeypatch)
    attempts = {"n": 0}

    def handler(request):
        if request.url.path == "/Tender/DownloadFile/119017" and attempts["n"] == 0:
            attempts["n"] += 1
            return httpx.Response(503)
        return base.handler(request)

    client = _client(httpx.MockTransport(handler))
    assert scraper.login(client)
    attachments, _, _ = scraper.download_documents(client, DOCS_ID, str(tmp_path))
    assert len(attachments) == 4 and attempts["n"] == 1


# ---------------------------------------------------------------------------
# One tender
# ---------------------------------------------------------------------------

def _entry(site_id=DETAIL_ID):
    return {"id": site_id, "url": f"{BASE}/OpenForBids/Details/{site_id}",
            "closes": "12/10/2026, at 05:00 PM", "opened": "26/09/2026"}


def test_scrape_opportunity_writes_the_text_file_with_the_full_page_and_source(tmp_path, monkeypatch):
    client, _ = _signed_in_client(monkeypatch)
    code, tender = scraper.scrape_opportunity(client, _entry(DOCS_ID), str(tmp_path), signed_in=True)
    folder = tmp_path / "TAS6649"
    text = (folder / "__tender__TAS6649.txt").read_text(encoding="utf-8")
    assert "TASMANIA ETENDERING:" in text
    assert common.FULL_PAGE_HEADER in text
    assert common.read_source_url(str(folder)) == f"{BASE}/OpenForBids/Details/{DOCS_ID}"
    assert tender["tender_id"] == "TAS6649" and len(tender["attachments"]) == 4


def test_the_page_text_is_saved_even_when_documents_are_skipped(tmp_path):
    code, tender = scraper.scrape_opportunity(_client(_portal()), _entry(), str(tmp_path), signed_in=False)
    assert code == common.TENDER_PARTIAL
    assert tender["documents_gated"] is True and tender["attachments"] == []
    assert (tmp_path / "TAS6652" / "__tender__TAS6652.txt").exists()


def test_a_tender_with_no_download_link_is_a_clean_success_without_logging_in(tmp_path):
    html = _detail().replace("/OpenForBids/Downloadables/", "/OpenForBids/Gone/")
    code, tender = scraper.scrape_opportunity(
        _client(_portal(detail_html=html)), _entry(), str(tmp_path), signed_in=False)
    assert code == common.SITE_SUCCESS and tender["documents_gated"] is False


def test_a_detail_page_that_cannot_be_parsed_is_a_structure_change(tmp_path):
    html = _detail().replace('class="itemrow', 'class="row-CHANGED')
    code, tender = scraper.scrape_opportunity(
        _client(_portal(detail_html=html)), _entry(), str(tmp_path), signed_in=True)
    assert code == common.SITE_STRUCTURE_CHANGE and tender == {}


def test_a_lapsed_session_while_signed_in_is_raised_with_the_tender_flagged(tmp_path, monkeypatch):
    client, _ = _signed_in_client(monkeypatch, lapse_after_logins=2)
    with pytest.raises(scraper.SessionLapsedError):
        scraper.scrape_opportunity(client, _entry(), str(tmp_path), signed_in=True)


def test_two_tenders_with_the_same_unique_id_get_separate_folders(tmp_path):
    # a portal that gives both tenders the same Unique Tender ID
    transport = httpx.MockTransport(lambda r: httpx.Response(200, text=_detail().replace(REFERENCE, "SAME-1")))
    client = _client(transport)
    taken = {}
    _, first = scraper.scrape_opportunity(client, _entry("1"), str(tmp_path), signed_in=False, taken=taken)
    _, second = scraper.scrape_opportunity(client, _entry("2"), str(tmp_path), signed_in=False, taken=taken)
    assert first["tender_id"] == "SAME-1" and second["tender_id"] == "SAME-1 (2)"
    assert first["folder"] != second["folder"]


def test_slow_tender_pages_and_files_are_given_two_minutes_or_more(tmp_path, monkeypatch):
    """A tender page that takes 40 s to answer used to time out at 30 s and cost the tender."""
    client, transport = _signed_in_client(monkeypatch)
    waits = {}
    real_send = client.send

    def spy(request, **kwargs):
        waits[request.url.path] = request.extensions["timeout"]["read"]
        return real_send(request, **kwargs)

    monkeypatch.setattr(client, "send", spy)
    scraper.scrape_opportunity(client, _entry(), str(tmp_path), signed_in=True)
    detail = [t for path, t in waits.items() if "Downloadables" not in path and "/Download" not in path]
    documents = [t for path, t in waits.items() if "Downloadables" in path]
    assert detail and min(detail) >= 120
    assert documents and min(documents) >= 120


def test_a_retried_tender_reuses_its_own_folder_name(tmp_path):
    client = _client(_portal())
    taken = {}
    _, first = scraper.scrape_opportunity(client, _entry(), str(tmp_path), signed_in=False, taken=taken)
    _, again = scraper.scrape_opportunity(client, _entry(), str(tmp_path), signed_in=False, taken=taken)
    assert first["tender_id"] == again["tender_id"] == "TAS6652"


# ---------------------------------------------------------------------------
# Whole run
# ---------------------------------------------------------------------------

def test_run_respects_limit_and_downloads_documents(tmp_path, monkeypatch):
    _set_credentials(monkeypatch)
    transport = _portal()
    _patch_client(monkeypatch, transport)
    result = scraper.run_scraper(limit=2, output_dir=str(tmp_path))
    # the fake xlsx is unreadable, so the tenders that fetch it are partial, but 2 are scraped
    assert result == common.ScrapeResult(common.TENDER_PARTIAL, "tenders_tas", 2)
    assert transport.state["logins"] == 1
    assert len(os.listdir(tmp_path)) == 2


def test_run_with_good_files_is_a_success(tmp_path, monkeypatch):
    _set_credentials(monkeypatch)
    good = dict(FILES)
    good["119019"] = ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", _xlsx_bytes)
    _patch_client(monkeypatch, _portal(files=good))
    result = scraper.run_scraper(limit=2, output_dir=str(tmp_path))
    assert result == common.ScrapeResult(common.SITE_SUCCESS, "tenders_tas", 2)


def test_run_with_a_wrong_password_is_login_failed_but_still_scrapes_the_text(tmp_path, monkeypatch):
    _set_credentials(monkeypatch, password="wrong")
    transport = _portal()
    _patch_client(monkeypatch, transport)
    result = scraper.run_scraper(limit=2, output_dir=str(tmp_path))
    assert result.code == common.SITE_LOGIN_FAILED and result.count == 2
    assert transport.state["failed_logins"] == 1
    assert (tmp_path / "TAS6652" / "__tender__TAS6652.txt").exists()


def test_run_without_credentials_scrapes_text_and_is_partial(tmp_path, monkeypatch):
    transport = _portal()
    _patch_client(monkeypatch, transport)
    result = scraper.run_scraper(limit=2, output_dir=str(tmp_path))
    assert result.code == common.TENDER_PARTIAL and result.count == 2
    assert not any(m == "POST" for m, _ in transport.state["requests"])


def test_run_signs_in_again_when_the_session_lapses(tmp_path, monkeypatch):
    _set_credentials(monkeypatch)
    transport = _portal(lapse_after_logins=2)
    _patch_client(monkeypatch, transport)
    result = scraper.run_scraper(limit=1, output_dir=str(tmp_path))
    assert transport.state["logins"] == 2
    assert result.count == 1 and len(os.listdir(tmp_path / "TAS6652")) > 2


def test_run_stops_re_logging_in_after_the_limit(tmp_path, monkeypatch):
    _set_credentials(monkeypatch)
    transport = _portal(lapse_after_logins=10_000)
    _patch_client(monkeypatch, transport)
    result = scraper.run_scraper(limit=1, output_dir=str(tmp_path))
    assert transport.state["logins"] <= scraper.MAX_RELOGINS + 1
    assert result.code == common.TENDER_PARTIAL


def test_empty_list_page_is_a_structure_change(tmp_path, monkeypatch):
    _patch_client(monkeypatch, _portal(list_html="<html><body>nothing</body></html>"))
    assert scraper.run_scraper(output_dir=str(tmp_path)).code == common.SITE_STRUCTURE_CHANGE


def test_site_total_failure_on_unreachable_url(tmp_path, monkeypatch):
    monkeypatch.setattr(common, "RUN_RETRY_WAITS", ())
    def handler(request):
        raise httpx.ConnectError("down")
    _patch_client(monkeypatch, httpx.MockTransport(handler))
    assert scraper.run_scraper(output_dir=str(tmp_path)).code == common.SITE_TOTAL_FAILURE


@pytest.mark.parametrize("status, expected", [
    (429, common.SITE_RATE_LIMITED),
    (403, common.SITE_BOT_BLOCKED),
])
def test_429_and_403_on_the_list_stop_the_run(status, expected, tmp_path, monkeypatch):
    calls = []

    def handler(request):
        calls.append(1)
        return httpx.Response(status)
    _patch_client(monkeypatch, httpx.MockTransport(handler))
    assert scraper.run_scraper(output_dir=str(tmp_path)).code == expected
    assert len(calls) == 1


@pytest.mark.parametrize("status, expected", [
    (429, common.SITE_RATE_LIMITED),
    (403, common.SITE_BOT_BLOCKED),
])
def test_429_and_403_on_a_detail_page_stop_the_run(status, expected, tmp_path, monkeypatch):
    base = _portal()

    def handler(request):
        if request.url.path.startswith("/OpenForBids/Details/"):
            return httpx.Response(status)
        return base.handler(request)
    _patch_client(monkeypatch, httpx.MockTransport(handler))
    result = scraper.run_scraper(output_dir=str(tmp_path))
    assert result.code == expected and result.count == 0


def test_a_detail_page_that_errors_makes_only_that_tender_partial(tmp_path, monkeypatch):
    base = _portal()

    def handler(request):
        if request.url.path == f"/OpenForBids/Details/{DETAIL_ID}":
            return httpx.Response(404)
        return base.handler(request)
    _patch_client(monkeypatch, httpx.MockTransport(handler))
    result = scraper.run_scraper(limit=2, output_dir=str(tmp_path))
    assert result.code == common.TENDER_PARTIAL and result.count == 1