#File holding our helper functions that are used across all scrapers. Also includes the status error codes for failures.

import logging
import multiprocessing
import os
import re
import time
from typing import NamedTuple
import fitz
import docx
import openpyxl
from docx.document import Document as _DocxDocument
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
import httpx

from error_scrapers import run_notes

#Error status codes for failures
#Scraping was a total success
SITE_SUCCESS = 0
#The URL provided couldnt be reached
SITE_TOTAL_FAILURE = 1
#The site login failed
SITE_LOGIN_FAILED = 2
#Anti bot detections have blocked us
SITE_BOT_BLOCKED = 3
#The structure of the website has changed thus causing our scraper to fail
SITE_STRUCTURE_CHANGE = 4
#Site has temporarily blocked our access
SITE_RATE_LIMITED = 5
#A tender only has partial information gathered and needs to be verified
TENDER_PARTIAL = 6

#-----
#This is what every scraper's run_scraper() returns -- and nothing else. The tenders
#themselves are left in the output folder for the pipeline to pick up.
#    code, site, count = run_scraper(...)      or      result.code / result.site / result.count
#-----
class ScrapeResult(NamedTuple):
    code: int    # one of the status codes above (0 = SITE_SUCCESS)
    site: str    # the portal's source_id, e.g. "grantconnect" -- same ids manager.py uses
    count: int   # how many tenders were scraped

#-----
#This is our folder and filename handling for the tenders
#-----

#This function is specifically for returning the folder path for a tender, and creating one if there isn't one.
#All of a tenders output, the page html, attachments, and extracted attachement text lives here and is named after the tender ID
def tender_dir(tender_id: str, base_dir: str = "tenders_data") -> str:
    run_notes.note_current(tender_id)  # if this tender fails, say which one
    path = os.path.join(base_dir, sanitise_filename(tender_id))
    os.makedirs(path, exist_ok=True)
    return path

#This function strips anything that could escape the tender folder or break the filesystem as we cant trust a name scraped off a page
def sanitise_filename(name: str) -> str:
    name = name.replace("/", "_").replace("\\", "_")
    name = re.sub(r'[<>:"|?*\x00-\x1f]', "_", name).strip()
    # "." and ".." are directories, not names: a tender ID of ".." would make
    # tender_dir() the folder above the output folder.
    return name if name not in ("", ".", "..") else "_"

#-----
#This is the section that saves the per tender output
#-----

#This function writes the scraped page content html content
def save_page_text(folder: str, tender_id: str, text: str) -> None:
    # Prefixed with __tender__ so the AI processor and attachment store can
    # distinguish this pipeline-generated file from real downloaded attachments.
    path = os.path.join(folder, f"__tender__{sanitise_filename(tender_id)}.txt")
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)

#-----
#Safety net for the page text: the scrapers only save the fields they have selectors for, so
#anything else on the details page (a site visit block, special instructions, a new section the
#portal adds) would never reach the AI. This keeps every visible line of the page as well.
#-----

FULL_PAGE_HEADER = "FULL PAGE TEXT (everything visible on the details page)"
FULL_PAGE_MAX_CHARS = 200_000

#Tags whose contents are never page text
_NOT_TEXT_TAGS = {"script", "style", "noscript", "head", "template", "svg", "iframe", "nav", "footer", "select", "option"}
#Tags that start a new line
_BLOCK_TAGS = {
    "address", "article", "aside", "blockquote", "body", "dd", "details", "div", "dl", "dt",
    "fieldset", "figcaption", "figure", "form", "h1", "h2", "h3", "h4", "h5", "h6", "header",
    "hr", "legend", "li", "main", "ol", "p", "pre", "section", "summary", "table", "tbody",
    "tfoot", "thead", "ul", "caption",
}


def visible_page_text(html: str, within: str | None = None) -> str:
    """All the visible text on a page, one line per block. Table rows stay on one line
    ('Label: value' for a two cell row, 'a | b | c' for wider ones). Scripts, styles, menus and
    footers are dropped; hidden panels are kept, since a portal can hide real content in a tab.
    `within` is a CSS selector for the page's content area (e.g. 'main'), which leaves out the
    site's menus and login panel; if it matches nothing the whole page is used."""
    from bs4 import BeautifulSoup, Comment, Doctype, NavigableString
    import html as html_lib

    def lines_of(root) -> list[str]:
        lines: list[str] = []
        buf: list[str] = []

        def flush():
            text = " ".join("".join(buf).split())
            buf.clear()
            if text and (not lines or lines[-1] != text):
                lines.append(text)

        def visit(node):
            if isinstance(node, (Comment, Doctype)):
                return
            if isinstance(node, NavigableString):
                buf.append(str(node))
                return
            name = node.name
            if name in _NOT_TEXT_TAGS:
                return
            if name == "br":
                flush()
                return
            if name == "tr":
                flush()
                cells = [" ".join(lines_of(cell)) for cell in node.find_all(["td", "th"], recursive=False)]
                cells = [cell for cell in cells if cell]
                if len(cells) == 2:
                    lines.append(f"{cells[0].rstrip(':')}: {cells[1]}")
                elif cells:
                    lines.append(" | ".join(cells))
                return
            if name == "textarea":
                flush()
                text = html_lib.unescape(node.get_text())
                if "<" in text:
                    text = "\n".join(lines_of(BeautifulSoup(text, "html.parser")))
                lines.extend(line.strip() for line in text.splitlines() if line.strip())
                return
            if name == "input":
                if node.get("type", "text") == "text" and node.get("value") and (
                        node.has_attr("readonly") or node.has_attr("disabled")):
                    buf.append(" " + node["value"] + " ")
                return
            block = name in _BLOCK_TAGS
            if block:
                flush()
            for child in node.children:
                visit(child)
            if block:
                flush()

        visit(root)
        flush()
        return lines

    soup = BeautifulSoup(html or "", "html.parser")
    root = (soup.select_one(within) if within else None) or soup
    try:
        lines = lines_of(root)
    except RecursionError:
        lines = [line for line in root.get_text("\n", strip=True).splitlines() if line.strip()]
    text = "\n".join(lines)
    if len(text) > FULL_PAGE_MAX_CHARS:
        text = text[:FULL_PAGE_MAX_CHARS] + "\n[page text cut at the size limit]"
    return text


def append_full_page_text(folder: str, tender_id: str, html: str, within: str | None = None) -> None:
    """Add the whole visible page to the tender's __tender__ text file, under a FULL PAGE TEXT
    heading, after whatever the scraper already wrote there (the file is created if it is not
    there yet). It never raises: the fields are already saved, and this is the extra."""
    try:
        text = visible_page_text(html, within)
        if not text:
            return
        path = os.path.join(folder, f"__tender__{sanitise_filename(tender_id)}.txt")
        with open(path, "a", encoding="utf-8") as f:
            f.write(f"\n{'=' * 80}\n{FULL_PAGE_HEADER}\n{'=' * 80}\n{text}\n")
    except Exception as exc:  # noqa: BLE001 - the safety net must never fail a tender
        logging.getLogger(__name__).warning("full page text not saved for %s (%s: %s)",
                                            tender_id, type(exc).__name__, exc)


SOURCE_URL_PATTERN = re.compile(r"Detail URL:\s*(https?://\S+)")


def add_source_url(folder: str, tender_id: str, url: str) -> None:
    """Record the tender's own page URL in its page-text file as a
    'Detail URL: <url>' line, so manager.py can read it back from the
    folder without the scraper having to return it."""
    path = os.path.join(folder, f"__tender__{sanitise_filename(tender_id)}.txt")
    with open(path, "a", encoding="utf-8") as f:
        f.write(f"\nDetail URL: {url}\n")


def read_source_url(folder: str) -> str | None:
    """The 'Detail URL:' a scraper recorded in this tender's page text, if any."""
    for name in sorted(os.listdir(folder)):
        if name.startswith("__tender__") and name.lower().endswith(".txt"):
            try:
                with open(os.path.join(folder, name), encoding="utf-8", errors="ignore") as f:
                    match = SOURCE_URL_PATTERN.search(f.read())
            except OSError:
                continue
            if match:
                return match.group(1).strip()
    return None


#This function makes sure we never silently overwrite a file. Portals often attach
#several files with the same name, so the second one would otherwise replace the first.
def unique_path(folder: str, filename: str) -> str:
    base, extension = os.path.splitext(sanitise_filename(filename))
    candidate = os.path.join(folder, f"{base}{extension}")
    counter = 2
    while os.path.exists(candidate):
        candidate = os.path.join(folder, f"{base} ({counter}){extension}")
        counter += 1
    return candidate

#This function focusses on the downloaded attachment from the tender page
def save_attachment(folder: str, filename: str, content: bytes) -> str:
    path = unique_path(folder, filename)
    with open(path, "wb") as f:
        f.write(content)
    return path

#Same as save_attachment, but writes an httpx streaming response straight to disk in
#chunks. Tender packs can include video, and the pipeline runs on Cloud Run where the
#filesystem is memory -- holding a whole file in RAM can take the run out.
def save_attachment_stream(folder: str, filename: str, response) -> str:
    path = unique_path(folder, filename)
    try:
        with open(path, "wb") as f:
            for chunk in response.iter_bytes():
                f.write(chunk)
    except BaseException:
        # A connection that drops mid-file must not leave half a file behind
        # (and a retry then writes the same name again).
        if os.path.exists(path):
            os.remove(path)
        raise
    return path

#This function is for extracted the attachment file's contents, we also make sure the file names match between the attachement and extracted information
def save_extracted_text(folder: str, attachment_filename: str, text: str) -> str:
    path = os.path.join(folder, f"{sanitise_filename(attachment_filename)}.txt")
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)
    return path

#-----
#These are the helper functions that help with the attachement information extraction
#-----

class StructureChangedError(Exception):
    """Raised when a page's response matches neither a known success or a known failure pattern."""

#This class is raised when the attachement cant be read for text
class ExtractionError(Exception):
    """Raised when a PDF/DOCX file can't be read for text."""


class EmptySourceFileError(ExtractionError):
    """
    Raised when a PDF has nothing in it at all (PyMuPDF: "cannot open empty
    document" / "Cannot open empty file"). The portal itself is serving a blank
    file, so there is no text to get and nothing for us to fix.
    """

class EncryptedFileError(ExtractionError):
    """
    Raised when a file is password protected. We cannot read it and will not be
    given the password, so the file is dropped (see extract_attachment_text) and
    the tender is reported as partial.
    """


#This is our pdf extraction code
def extract_pdf(file_path: str) -> str:
    try:
        text = ""
        with fitz.open(file_path) as pdf_doc:
            for page in pdf_doc:
                text += page.get_text("text") + "\n"
        return text
    except Exception as e:
        if "empty" in str(e).lower():
            raise EmptySourceFileError(f"PDF extraction failed on {file_path}: {e}") from e
        raise ExtractionError(f"PDF extraction failed on {file_path}: {e}") from e

def _iter_block_items(parent):
    """
    Yield each paragraph and table in a docx document (or table cell),
    in the order they actually appear -- not paragraphs-then-tables,
    but interleaved as a person reading top to bottom would encounter
    them. Walks the underlying XML body directly, since python-docx's
    own doc.paragraphs and doc.tables are two separate, unordered
    collections with no way to tell which came first.
    """
    if isinstance(parent, _DocxDocument):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("_iter_block_items: unsupported parent type")
 
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)
 
 
def _element_text(element):
    """
    Every visible text node inside an XML element, regardless of
    whether it sits in a plain run, a hyperlink, a content control /
    dropdown (w:sdt), or a floating textbox / watermark (w:drawing or
    the legacy w:pict). Scanning all descendant w:t nodes is
    structure-agnostic, unlike the higher-level .text properties
    (paragraph.text, cell.text), which only walk certain expected
    shapes and miss anything nested one level deeper than that.
    """
    return "".join(t.text or "" for t in element.iter(qn("w:t")))
 
 
def _table_to_text(table):
    """Render one table's cells as readable, tab-separated rows."""
    lines = []
    # Walk each row's own <w:tc> cells. python-docx's row.cells rebuilds the
    # whole table grid for every row (quadratic on big tables) and raises
    # IndexError on irregular merged cells.
    for tr in table._tbl.tr_lst:
        cells = [_element_text(tc) for tc in tr.tc_lst]
        if any(cell.strip() for cell in cells):  # skip fully-empty rows
            lines.append("\t".join(cells))
    return "\n".join(lines)
 
 
def _all_header_footer_elements(document):
    """
    Every header/footer XML part actually referenced by the document,
    including 'even page' headers/footers -- which python-docx's public
    API does not expose at all. Reads each section's raw
    headerReference/footerReference elements and resolves them through
    the document part's own relationships, rather than trusting
    section.header/.footer/.first_page_header/.first_page_footer, which
    can silently return the wrong (or a newly-created, empty) part when
    a document uses header/footer types those properties don't know
    about.
    """
    doc_part = document.part
    seen_rids, elements = set(), []
    for section in document.sections:
        sect_pr = section._sectPr
        for tag in (qn("w:headerReference"), qn("w:footerReference")):
            for ref in sect_pr.findall(tag):
                r_id = ref.get(qn("r:id"))
                if r_id and r_id not in seen_rids and r_id in doc_part.rels:
                    seen_rids.add(r_id)
                    elements.append(doc_part.rels[r_id].target_part.element)
    return elements
 
 
def extract_docx(file_path: str) -> str:
    """
    Extract all readable text from a Word document: headers, footers
    (including even-page variants and watermarks), body paragraphs,
    tables, and content-control/dropdown text -- everything a person
    would actually see when reading the document, not just its main
    paragraph flow.
    """
    try:
        doc = docx.Document(file_path)
        parts = []
 
        for element in _all_header_footer_elements(doc):
            text = _element_text(element)
            if text.strip():
                parts.append(text)
 
        for block in _iter_block_items(doc):
            if isinstance(block, Paragraph):
                text = _element_text(block._p)
            elif isinstance(block, Table):
                text = _table_to_text(block)
            else:
                continue
            if text.strip():
                parts.append(text)
 
        return "\n".join(part for part in parts if part.strip())
    except Exception as e:
        raise ExtractionError(f"DOCX extraction failed on {file_path}: {e}") from e


# A Word file's main text part is never legitimately this large.
_MAX_DOCX_XML_BYTES = 200 * 1024 * 1024


def extract_docx_body_text(file_path: str) -> str:
    """
    Last-resort DOCX text: read the body straight out of the file's zip, one
    line per paragraph, without python-docx. Used when the full extraction
    fails or runs out of time (a package python-docx refuses to open, a table
    that makes it crawl). Table cell text comes out as ordinary lines, and
    headers and footers are left out.
    """
    import zipfile
    from lxml import etree

    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    try:
        with zipfile.ZipFile(file_path) as archive:
            if archive.getinfo("word/document.xml").file_size > _MAX_DOCX_XML_BYTES:
                raise ValueError("document.xml is implausibly large")
            xml = archive.read("word/document.xml")
        parser = etree.XMLParser(resolve_entities=False, no_network=True)
        root = etree.fromstring(xml, parser)
        lines = []
        for paragraph in root.iter(ns + "p"):
            text = "".join(t.text or "" for t in paragraph.iter(ns + "t"))
            if text.strip():
                lines.append(text)
        return "\n".join(lines)
    except Exception as e:
        raise ExtractionError(f"DOCX body-text extraction failed on {file_path}: {e}") from e

def extract_xlsx(file_path: str) -> str:
    """
    Extract all cell content from every sheet in a spreadsheet.
 
    data_only=True returns each formula cell's last-calculated value
    rather than the formula text itself -- what a person looking at the
    spreadsheet would actually see. read_only=True keeps memory use low
    on large files.
    """
    try:
        wb = load_workbook(file_path, data_only=True, read_only=True)
        parts = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            # Some files declare a sheet size of A1:XFD1048576 (stray
            # formatting), which makes every row come back 16,384 cells wide.
            # Ignoring the declared size reads only the cells actually stored.
            ws.reset_dimensions()
            sheet_lines = [f"## Sheet: {sheet_name}"]
            for row in ws.iter_rows(values_only=True):
                cells = ["" if cell is None else str(cell) for cell in row]
                while cells and not cells[-1].strip():  # drop trailing empty cells
                    cells.pop()
                if any(cell.strip() for cell in cells):  # skip fully-empty rows
                    sheet_lines.append("\t".join(cells))
            if len(sheet_lines) > 1:  # only keep sheets that had real content
                parts.append("\n".join(sheet_lines))
        wb.close()
        return "\n\n".join(parts)
    except Exception as e:
        raise ExtractionError(f"XLSX extraction failed on {file_path}: {e}") from e

# A spreadsheet's XML parts are never legitimately this large.
_MAX_XLSX_XML_BYTES = 1024 * 1024 * 1024
_SHEET_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_PKG_REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def _column_index(cell_ref: str) -> int:
    """0-based column of a cell reference such as "C12" (AA7 is 26)."""
    index = 0
    for char in cell_ref:
        if not char.isalpha():
            break
        index = index * 26 + (ord(char.upper()) - 64)
    return max(index - 1, 0)


def extract_xlsx_cells(file_path: str) -> str:
    """
    Fallback spreadsheet text: read the cell values straight out of the file's
    zip, in the same layout extract_xlsx produces, without openpyxl. Used when
    openpyxl refuses a workbook over something that has nothing to do with its
    cells -- for example a cell style with no name, which some estimating
    software writes (openpyxl: "_NamedCellStyle.name should be <class 'str'>").

    Cached values only (a formula shows its last-calculated result), and dates
    come out as their serial numbers.
    """
    import zipfile
    from lxml import etree

    def parse(archive, name):
        if archive.getinfo(name).file_size > _MAX_XLSX_XML_BYTES:
            raise ValueError(f"{name} is implausibly large")
        parser = etree.XMLParser(resolve_entities=False, no_network=True)
        return etree.fromstring(archive.read(name), parser)

    try:
        with zipfile.ZipFile(file_path) as archive:
            names = set(archive.namelist())

            shared: list[str] = []
            if "xl/sharedStrings.xml" in names:
                for item in parse(archive, "xl/sharedStrings.xml").iter(_SHEET_NS + "si"):
                    # Rich text is several runs; phonetic hints (rPh) are not part of the text.
                    shared.append("".join(
                        t.text or "" for t in item.iter(_SHEET_NS + "t")
                        if not any(parent.tag == _SHEET_NS + "rPh" for parent in t.iterancestors())))

            sheets: list[tuple[str, str]] = []  # (sheet name, part name)
            try:
                targets = {rel.get("Id"): rel.get("Target") for rel in
                           parse(archive, "xl/_rels/workbook.xml.rels").iter(_PKG_REL_NS + "Relationship")}
                for sheet in parse(archive, "xl/workbook.xml").iter(_SHEET_NS + "sheet"):
                    target = targets.get(sheet.get(_REL_NS + "id"))
                    if target:
                        part = target.lstrip("/") if target.startswith("/") else "xl/" + target
                        if part in names:
                            sheets.append((sheet.get("name") or part, part))
            except Exception:
                sheets = []
            if not sheets:
                sheets = [(f"Sheet{n}", name) for n, name in enumerate(
                    sorted(n for n in names if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n)), start=1)]

            parts = []
            for sheet_name, part in sheets:
                lines = [f"## Sheet: {sheet_name}"]
                for row in parse(archive, part).iter(_SHEET_NS + "row"):
                    cells: dict[int, str] = {}
                    for position, cell in enumerate(row.iter(_SHEET_NS + "c")):
                        column = _column_index(cell.get("r") or "") if cell.get("r") else position
                        kind = cell.get("t")
                        value = cell.find(_SHEET_NS + "v")
                        if kind == "inlineStr":
                            text = "".join(t.text or "" for t in cell.iter(_SHEET_NS + "t"))
                        elif value is None or value.text is None:
                            continue
                        elif kind == "s":
                            index = int(value.text)
                            text = shared[index] if 0 <= index < len(shared) else ""
                        elif kind == "b":
                            text = "True" if value.text.strip() == "1" else "False"
                        else:
                            text = value.text
                        cells[column] = text
                    if not cells:
                        continue
                    line = ["" for _ in range(max(cells) + 1)]
                    for column, text in cells.items():
                        line[column] = text
                    while line and not line[-1].strip():
                        line.pop()
                    if any(text.strip() for text in line):
                        lines.append("\t".join(line))
                if len(lines) > 1:
                    parts.append("\n".join(lines))
            return "\n\n".join(parts)
    except Exception as e:
        raise ExtractionError(f"XLSX cell-value extraction failed on {file_path}: {e}") from e


def extract_xls(file_path: str) -> str:
    """
    Extract all cell content from every sheet of a legacy Excel (.xls)
    workbook, in the same layout extract_xlsx produces. Agencies still
    publish schedules of rates in this pre-2007 format.
    """
    try:
        import xlrd

        book = xlrd.open_workbook(file_path, on_demand=True)
        parts = []
        for sheet in book.sheets():
            sheet_lines = [f"## Sheet: {sheet.name}"]
            for row_index in range(sheet.nrows):
                cells = [str(cell.value) if cell.value is not None else ""
                         for cell in sheet.row(row_index)]
                cells = [c[:-2] if c.endswith(".0") and c[:-2].lstrip("-").isdigit() else c
                         for c in cells]
                if any(cell.strip() for cell in cells):
                    sheet_lines.append("\t".join(cells))
            if len(sheet_lines) > 1:
                parts.append("\n".join(sheet_lines))
        book.release_resources()
        return "\n\n".join(parts)
    except Exception as e:
        if "encrypted" in str(e).lower():
            raise EncryptedFileError(f"XLS extraction failed on {file_path}: {e}") from e
        raise ExtractionError(f"XLS extraction failed on {file_path}: {e}") from e


def extract_pptx(file_path: str) -> str:
    """Extract the text of every slide (shapes, tables and speaker notes)."""
    try:
        from pptx import Presentation

        parts = []
        for number, slide in enumerate(Presentation(file_path).slides, start=1):
            lines = [f"## Slide {number}"]
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    lines.append(shape.text_frame.text.strip())
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        lines.append("\t".join(cell.text for cell in row.cells))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                lines.append("Notes: " + slide.notes_slide.notes_text_frame.text.strip())
            if len(lines) > 1:
                parts.append("\n".join(lines))
        return "\n\n".join(parts)
    except Exception as e:
        raise ExtractionError(f"PPTX extraction failed on {file_path}: {e}") from e


def extract_plain_text(file_path: str) -> str:
    """Read a plain-text or CSV file, tolerating the odd Windows-1252 file."""
    try:
        with open(file_path, "rb") as f:
            raw = f.read()
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                return raw.decode(encoding)
            except UnicodeDecodeError:
                continue
        return raw.decode("utf-8", errors="replace")
    except Exception as e:
        raise ExtractionError(f"Text extraction failed on {file_path}: {e}") from e


#The file types we can pull text out of. Anything else (images, video, legacy
#.doc/.ppt) is still saved as an attachment, it just has no extracted .txt next to it.
EXTRACTORS = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".xlsx": extract_xlsx,
    ".xlsm": extract_xlsx,
    ".xls": extract_xls,
    ".pptx": extract_pptx,
    ".csv": extract_plain_text,
    ".txt": extract_plain_text,
}

# One unreadable file must not stall a whole night's run: each extraction gets
# this long, and is stopped (and reported as a failure) if it takes longer.
EXTRACTION_TIMEOUT_SECONDS = 300

logger = logging.getLogger("scraper.common")


def _extract_in_child(extractor, path, conn):
    try:
        conn.send(("ok", extractor(path)))
    except Exception as e:  # reported to the parent, which raises ExtractionError
        conn.send(("error", f"{type(e).__name__}: {e}"))
    finally:
        conn.close()


def run_extractor(extractor, path: str, timeout: float | None = None) -> str:
    """
    Run one text extractor with a time limit. The work happens in a separate
    process so a stuck file can actually be killed. Raises ExtractionError if
    the extractor fails or runs out of time. Where processes can't be forked
    (Windows), the extractor runs in-process with no limit.
    """
    timeout = EXTRACTION_TIMEOUT_SECONDS if timeout is None else timeout
    if "fork" not in multiprocessing.get_all_start_methods():
        return extractor(path)
    ctx = multiprocessing.get_context("fork")
    parent_conn, child_conn = ctx.Pipe(duplex=False)
    proc = ctx.Process(target=_extract_in_child, args=(extractor, path, child_conn), daemon=True)
    proc.start()
    child_conn.close()
    try:
        if not parent_conn.poll(timeout):
            raise ExtractionError(
                f"timed out after {timeout:.0f}s on {os.path.basename(path)}; skipped its text"
            )
        try:
            status, payload = parent_conn.recv()
        except EOFError:
            raise ExtractionError(
                f"extractor died without a result on {os.path.basename(path)}"
            ) from None
    finally:
        if proc.is_alive():
            proc.kill()
        proc.join()
        parent_conn.close()
    if status == "error":
        # The type name travels in the message, because only text crosses the process boundary.
        if payload.startswith("EmptySourceFileError:"):
            raise EmptySourceFileError(payload)
        if payload.startswith("EncryptedFileError:"):
            raise EncryptedFileError(payload)
        raise ExtractionError(payload)
    return payload


_ZIP_OFFICE_TYPES = (".docx", ".xlsx", ".xlsm", ".pptx")
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
# "EncryptedPackage" as it is stored inside an encrypted Office file (UTF-16).
_ENCRYPTED_MARKER = "EncryptedPackage".encode("utf-16-le")


def _check_office_container(path: str):
    """
    Look at the first bytes of a .docx/.xlsx/.xlsm/.pptx before handing it to an
    extractor. These formats are zip files; anything else is not what its name
    says.

    Returns None for a genuine zip (carry on as normal). Returns the file's text
    ("" if it has none we can read) for a file that is not a zip, after logging
    what it actually is (.docx only; other types fail as before). Raises EncryptedFileError for a password-protected one.
    """
    name = os.path.basename(path)
    is_docx = path.lower().endswith(".docx")
    try:
        with open(path, "rb") as handle:
            head = handle.read(8)
            if head.startswith(b"PK"):
                return None
            if head == _OLE_MAGIC:
                handle.seek(0)
                if _ENCRYPTED_MARKER in handle.read():
                    raise EncryptedFileError(
                        f"{name} is password protected (encrypted Office file)")
                if not is_docx:
                    return None
                kind = "an old-format Office file (not a zip)"
                text = ""
            elif not is_docx:
                return None  # other types: the extractor fails as before
            else:
                handle.seek(0)
                sample = handle.read(5 * 1024 * 1024)
                stripped = sample.lstrip()[:200].lower()
                if stripped.startswith((b"<!doctype html", b"<html", b"<?xml", b"<")):
                    kind = "a web/XML page"
                    text = re.sub(r"\s+", " ", re.sub(
                        r"<[^>]*>", " ", sample.decode("utf-8", errors="replace"))).strip()
                elif not sample:
                    kind = "an empty file"
                    text = ""
                else:
                    kind = "not a zip file (unrecognised content)"
                    text = ""
    except OSError:
        return None  # let the normal extractor report the real problem
    logger.warning("%s is named %s but is %s; stored %s",
                   name, os.path.splitext(name)[1].lower(), kind,
                   "with its readable text" if text else "without text")
    return text


def extract_with_fallback(extractor, path: str, timeout: float | None = None) -> str:
    """
    run_extractor, plus one more try for Word and Excel files: if the full
    extraction fails or times out, fall back to the text read straight from the
    file (a Word file's body text; a spreadsheet's cell values). The tender
    keeps that text (tables, headers, formatting may be missing), and the loss
    is logged. Any other file type, or a file the fallback can't read either,
    raises the original error.

    A password-protected file raises EncryptedFileError. An Office-named file
    that is not a zip at all (a web page or old-format file saved under the
    wrong name) is stored with whatever text it has, and logged -- the portal
    served it that way, so there is nothing for us to fix.
    """
    if path.lower().endswith(_ZIP_OFFICE_TYPES):
        text = _check_office_container(path)
        if text is not None:
            return text
    try:
        return run_extractor(extractor, path, timeout)
    except EmptySourceFileError as e:
        logger.warning("%s is empty at the source (%s); stored without text",
                       os.path.basename(path), e)
        return ""
    except ExtractionError as first_error:
        lower = path.lower()
        if lower.endswith(".docx"):
            fallback, kept = extract_docx_body_text, "its body text only"
        elif lower.endswith((".xlsx", ".xlsm")):
            fallback, kept = extract_xlsx_cells, "its cell values only"
        else:
            raise
        if extractor is fallback:
            raise
        try:
            text = run_extractor(fallback, path, timeout)
        except ExtractionError:
            raise first_error from None
        logger.warning("full extraction failed for %s (%s); kept %s",
                       os.path.basename(path), first_error, kept)
        return text


def _drop_file(path: str) -> None:
    for candidate in (path, path + ".txt"):
        try:
            os.remove(candidate)
        except OSError:
            pass


def drop_locked_file(path: str) -> None:
    """
    A password-protected file we cannot read and will never be given the
    password for: remove it (and any text file beside it), log it, and note the
    reason against the tender being scraped.
    """
    name = os.path.basename(path)
    _drop_file(path)
    logger.warning("%s is password protected; dropped it (not stored)", name)
    run_notes.note_reason(run_notes.LOCKED, name)


def zip_member_is_encrypted(member) -> bool:
    """True if a zip entry is password protected (the zip's own encryption flag)."""
    return bool(member.flag_bits & 0x1)


def extract_attachment_text(folder: str, file_name: str) -> bool:
    """
    Write <file_name>.txt next to a saved attachment, when we have an extractor
    for its type. Returns False only when extraction of a supported type
    failed -- an unsupported type is not a failure, there's just no text in it.
    """
    extractor = EXTRACTORS.get(os.path.splitext(file_name)[1].lower())
    if extractor is None:
        return True
    logger.info(f"extracting text from {file_name}")
    try:
        path = os.path.join(folder, file_name)
        text = extract_with_fallback(extractor, path)
        save_extracted_text(folder, file_name, text)
        return True
    except EncryptedFileError as e:
        # Password protected: we can't read it and never will. Drop the file so
        # it isn't stored, say so in the log, and report the tender as partial.
        _drop_file(path)
        print(f"PASSWORD PROTECTED: {file_name} dropped")
        logger.warning("%s is password protected; dropped it (not stored) -- "
                       "tender %s is reported as partial", file_name,
                       os.path.basename(os.path.normpath(folder)))
        run_notes.note_reason(run_notes.LOCKED, file_name)
        return False
    except ExtractionError as e:
        size = os.path.getsize(path) if os.path.exists(path) else "?"
        print(f"EXTRACTION FAILED: {e}")
        logger.warning(f"text extraction failed for {file_name} ({size} bytes): {e}")
        run_notes.note_reason(run_notes.UNREADABLE, file_name)
        return False

# A page request that fails in a way that usually passes on its own is made
# again, this many attempts in all, waiting a little longer each time:
#   * a timeout, a dropped or refused connection, a connection cut mid-reply
#   * a 500, 502, 503 or 504 from the server
# Never a 403 or 429 (the portal telling us to stop), never anything the site
# sends as a block or challenge, and a login is never repeated (a lockout risk).
PAGE_ATTEMPTS = 3
PAGE_RETRY_SECONDS = 5
RETRY_STATUSES = frozenset({500, 502, 503, 504})


def is_block_response(response) -> bool:
    """True if the site sent this as a block or challenge (Cloudflare, AWS WAF)."""
    headers = response.headers
    return headers.get("cf-mitigated", "").lower() == "challenge" or "x-amzn-waf-action" in headers


def is_transient_error(exc: BaseException) -> bool:
    """
    True for a failure that is usually gone on the next try: a timeout, a
    dropped or refused connection, a connection cut mid-reply, or a 500/502/503/504.
    False for everything else, including a 403, 404 or 429 and anything the
    site sent as a block.
    """
    if isinstance(exc, httpx.HTTPStatusError):
        return exc.response.status_code in RETRY_STATUSES and not is_block_response(exc.response)
    return isinstance(exc, (httpx.TimeoutException, httpx.NetworkError, httpx.RemoteProtocolError))


def request_with_retries(client, method: str, url: str, **kwargs):
    """
    client.get / client.post (`method` is "get" or "post"), asked again when the
    request times out, the connection drops, or the server answers 500/502/503/504.

    If every attempt fails the last error is raised, or for a server error the
    last response is returned so the caller's raise_for_status sees it. Any
    other status (403, 404, 429...) comes straight back, never retried.
    """
    send = getattr(client, method)
    log = logging.getLogger("scraper.common")
    for attempt in range(1, PAGE_ATTEMPTS + 1):
        try:
            response = send(url, **kwargs)
        except (httpx.TimeoutException, httpx.NetworkError, httpx.RemoteProtocolError) as exc:
            if attempt == PAGE_ATTEMPTS:
                raise
            problem = type(exc).__name__
        else:
            if (response.status_code not in RETRY_STATUSES or is_block_response(response)
                    or attempt == PAGE_ATTEMPTS):
                return response
            problem = f"HTTP {response.status_code}"
            response.close()
        wait = PAGE_RETRY_SECONDS * attempt
        log.warning("       %s failed (%s) -- trying again in %ds (attempt %d/%d)",
                    url, problem, wait, attempt + 1, PAGE_ATTEMPTS)
        time.sleep(wait)


# A document download that fails the same way is made again once. Used by every
# scraper that streams a file (GrantConnect, Austender, QLD, NT).
DOWNLOAD_ATTEMPTS = 2
DOWNLOAD_RETRY_SECONDS = 5


def retry_transient(action, what: str, attempts: int | None = None):
    """
    Run `action()` (which fetches one document) and run it again if it fails
    with a transient error (see is_transient_error). Anything else, such as a
    login page served in place of the file, is raised at once.
    """
    attempts = DOWNLOAD_ATTEMPTS if attempts is None else attempts
    for attempt in range(1, attempts + 1):
        try:
            return action()
        except Exception as exc:
            if attempt == attempts or not is_transient_error(exc):
                raise
            wait = DOWNLOAD_RETRY_SECONDS * attempt
            logging.getLogger("scraper.common").warning(
                "       %s failed (%s: %s) -- trying again in %ds (attempt %d/%d)",
                what, type(exc).__name__, exc, wait, attempt + 1, attempts)
            time.sleep(wait)


# A browser page that fails to LOAD (the driver times out, the connection is
# reset) is loaded again, this many attempts in all, waiting longer each time.
# Only the load is retried: a Cloudflare challenge, a login page or a missing
# element is not a load failure and is never reloaded.
PAGE_LOAD_ATTEMPTS = 3
PAGE_LOAD_RETRY_SECONDS = 5
_TRANSIENT_BROWSER_ERRORS = (
    "ERR_CONNECTION_RESET", "ERR_CONNECTION_CLOSED", "ERR_CONNECTION_TIMED_OUT",
    "ERR_TIMED_OUT", "ERR_EMPTY_RESPONSE", "ERR_NETWORK_CHANGED", "ERR_CONNECTION_REFUSED",
)


def is_transient_page_load_error(exc: BaseException) -> bool:
    """
    True for a browser page load that failed in a way a reload usually fixes: a
    timeout from the driver (selenium's TimeoutException, urllib3's
    ReadTimeoutError, TimeoutError) or Chrome reporting a reset, closed, empty
    or timed-out connection. Nothing else, so a dead browser still fails at once.
    """
    if isinstance(exc, TimeoutError) or "timeout" in type(exc).__name__.lower():
        return True
    return any(code in str(exc) for code in _TRANSIENT_BROWSER_ERRORS)


def load_page_with_retries(load, url: str):
    """Run `load()` (which opens `url` in the browser), loading again if the load itself fails."""
    for attempt in range(1, PAGE_LOAD_ATTEMPTS + 1):
        try:
            return load()
        except Exception as exc:
            if attempt == PAGE_LOAD_ATTEMPTS or not is_transient_page_load_error(exc):
                raise
            wait = PAGE_LOAD_RETRY_SECONDS * attempt
            logging.getLogger("scraper.common").warning(
                "       %s did not load (%s) -- loading it again in %ds (attempt %d/%d)",
                url, type(exc).__name__, wait, attempt + 1, PAGE_LOAD_ATTEMPTS)
            time.sleep(wait)


#Raised when a portal answers a document request with an HTML page (almost
#always its login form) instead of the file.
class DocumentGatedError(Exception):
    """Raised when a document download comes back as a web page, not a file."""

def stream_attachment(client, url: str, folder: str, file_name: str,
                      headers: dict | None = None, timeout: float = 60.0,
                      reject_web_pages: bool = True) -> tuple[str, str | None]:
    """
    Stream one file into `folder` and return (path written, content type). A
    timeout, dropped connection or 500/502/503/504 is tried once more
    (retry_transient); the half-written file is removed first.

    With reject_web_pages, a web page served in place of a non-.htm file (almost
    always the portal's login form) raises DocumentGatedError, which is not retried.
    """
    def fetch_once():
        with client.stream("GET", url, headers=headers or {}, timeout=timeout) as response:
            response.raise_for_status()
            content_type = response.headers.get("content-type")
            wants_html = file_name.lower().endswith((".htm", ".html"))
            if (reject_web_pages and content_type and "text/html" in content_type.lower()
                    and not wants_html):
                raise DocumentGatedError(f"{url} returned a web page, not a file")
            return save_attachment_stream(folder, file_name, response), content_type

    return retry_transient(fetch_once, url)


def download_attachment(client, url: str, folder: str, file_name: str,
                        headers: dict | None = None, timeout: float = 60.0) -> tuple[dict, bool]:
    """
    Stream one attachment into the tender folder and extract its text.

    Returns (attachment, extracted). `attachment` is the manifest entry the
    pipeline needs ({file_name, content_type, size_bytes}), using the name
    actually written -- save_attachment_stream may have de-duplicated it.
    `extracted` is False when text extraction failed on a supported type, so
    the caller can report the tender as partial.

    Raises on a failed download (HTTP error, network error, or a login page
    served in place of the file), so the caller decides what that means. A
    transient failure is tried once more first (stream_attachment).
    """
    raw_path, content_type = stream_attachment(client, url, folder, file_name, headers, timeout)

    saved_name = os.path.basename(raw_path)
    if os.path.getsize(raw_path) == 0:
        os.remove(raw_path)
        raise ExtractionError(f"{url} returned an empty file")

    attachment = {
        "file_name": saved_name,
        "content_type": content_type,
        "size_bytes": os.path.getsize(raw_path),
    }
    extracted = extract_attachment_text(folder, saved_name)
    if not os.path.exists(raw_path):
        # Dropped because it is password protected: it is not an attachment.
        raise EncryptedFileError(f"{saved_name} is password protected; dropped")
    return attachment, extracted

def wait_for_download(downloads_dir: str, before, max_wait: float,
                      poll: float = 1.0, stall_seconds: float = 300.0,
                      hard_max: float = 900.0, should_stop=None,
                      stop_check_seconds: float = 5.0) -> str | None:
    """
    Wait for the browser to finish saving one new file into `downloads_dir`.

    `before` is the set of file names that were already there. Returns the new
    file's name, or None if nothing finished.

    Two stages. First, the download has to *start*: some new file, finished or
    in progress, must show up within `max_wait` seconds. Chrome writes an
    in-progress download as "<name>.crdownload". Once one exists the portal has
    answered, so it is given time: the wait continues until the file finishes,
    giving up only if the partial file has not changed in `stall_seconds`, or
    `hard_max` seconds have passed in all. A big package is slow, not broken,
    and a partial file that is still alive is not abandoned.

    `should_stop`, if given, is called every `stop_check_seconds` while no file has
    shown up yet; when it returns True the wait ends at once with None. A browser
    that is looking at a block page will never produce a file, so there is no
    point waiting the full `max_wait` for one.
    """
    before = set(before)
    # A download the previous tender gave up on (its "X.zip.crdownload" was
    # already there) can finish during this wait, under a name that is not in
    # `before`. It is that tender's file, not this one's, so it is not "new".
    stale = {name[: -len(".crdownload")] for name in before if name.endswith(".crdownload")}
    started = time.monotonic()
    last_change = started
    last_stop_check = started
    seen_sizes: dict[str, int] = {}
    seen_partial = False
    vanished_polls = 0
    while True:
        new = set(os.listdir(downloads_dir)) - before - stale
        finished = sorted(name for name in new if not name.endswith(".crdownload"))
        if finished:
            return finished[0]
        now = time.monotonic()
        partials = sorted(new)
        if partials:
            seen_partial = True
            vanished_polls = 0
            for name in partials:
                try:
                    size = os.path.getsize(os.path.join(downloads_dir, name))
                except OSError:
                    continue
                if seen_sizes.get(name) != size:
                    seen_sizes[name] = size
                    last_change = now
            if now - last_change > stall_seconds or now - started > hard_max:
                return None
        else:
            if should_stop is not None and not seen_partial and now - last_stop_check >= stop_check_seconds:
                last_stop_check = now
                try:
                    if should_stop():
                        return None
                except Exception:
                    pass  # a check that cannot run must not end the wait
            if seen_partial:
                # An in-progress file disappeared without a finished one: cancelled.
                vanished_polls += 1
                if vanished_polls >= 3:
                    return None
            elif now - started > max_wait:
                return None
        time.sleep(poll)


# Files that ride along in a package but are never documents: Office lock files
# ("~$report.pdf", 162 bytes of who had it open), macOS resource-fork copies and
# Windows folder metadata. They are skipped, not stored, and are not a failure.
JUNK_FILE_NAMES = frozenset({".ds_store", "thumbs.db", "desktop.ini"})
# An empty file of one of these kinds is leftover system data (e.g. the empty
# "TRIMfiles.dat" that TRIM/Content Manager exports leave behind), not a
# download that went wrong. An empty .pdf or .docx still is.
EMPTY_JUNK_EXTENSIONS = frozenset({"", ".dat", ".tmp", ".db", ".ini", ".bak", ".log", ".lnk"})


def is_junk_file(member_name: str) -> bool:
    """True for a package member that is never a document (see JUNK_FILE_NAMES)."""
    path = member_name.replace("\\", "/")
    name = os.path.basename(path).lower()
    return (name.startswith("~$") or name.startswith("._") or name in JUNK_FILE_NAMES
            or "__macosx/" in path.lower())


MAX_ZIP_DEPTH = 5
# A package member larger than this is not unpacked. attachment_store would
# refuse to store it anyway (MAX_ATTACHMENT_MB), and on Cloud Run the
# filesystem is memory: writing it out could take the run down first.
MAX_PACKAGE_MEMBER_BYTES = 500 * 1024 * 1024

def unpack_zip(zip_file, folder: str, depth: int = 0) -> tuple[list[dict], bool]:
    """
    Unpack a downloaded document bundle into the tender folder, extracting
    text from each file. `zip_file` is a path or a file-like object.

    Nested zips (agencies bundle response templates this way) are recursed
    into rather than saved as an unreadable blob. Member names are reduced to
    a bare filename, so an archive holding ../../evil.txt lands inside the
    tender folder or not at all.

    Returns (attachments, any_failed).
    """
    import io
    import zipfile

    attachments, any_failed = [], False
    if depth > MAX_ZIP_DEPTH:
        return attachments, True

    with zipfile.ZipFile(zip_file) as zf:
        for member in zf.infolist():
            file_name = os.path.basename(member.filename.replace("\\", "/"))
            if member.is_dir() or not file_name:
                continue
            if is_junk_file(member.filename):
                logger.info("package held a junk file (%s); ignored", file_name)
                continue
            if zip_member_is_encrypted(member):
                # One locked file must not cost the tender the rest of its package.
                logger.warning("package held a password-protected file (%s); not stored", file_name)
                run_notes.note_reason(run_notes.LOCKED, file_name)
                any_failed = True
                continue
            if member.file_size > MAX_PACKAGE_MEMBER_BYTES:
                logger.warning("package held a file over the size limit (%s, %d MB); not stored",
                               file_name, member.file_size // (1024 * 1024))
                run_notes.note_reason(run_notes.SIZE, file_name)
                any_failed = True
                continue
            if file_name.lower().endswith(".zip"):
                nested, nested_failed = unpack_zip(io.BytesIO(zf.read(member)), folder, depth + 1)
                attachments += nested
                any_failed = any_failed or nested_failed
                continue
            with zf.open(member) as source:
                path = unique_path(folder, file_name)
                with open(path, "wb") as target:
                    while chunk := source.read(1024 * 1024):
                        target.write(chunk)
            saved_name = os.path.basename(path)
            if os.path.getsize(path) == 0 and os.path.splitext(saved_name)[1].lower() in EMPTY_JUNK_EXTENSIONS:
                os.remove(path)
                logger.info("package held an empty %s file (%s); ignored",
                            os.path.splitext(saved_name)[1].lower() or "extensionless", saved_name)
                continue
            if os.path.getsize(path) == 0:
                # An empty document in the package is a download that went wrong,
                # not a document: drop it and report the tender as partial.
                os.remove(path)
                logger.warning(f"package held an empty file ({saved_name}); not stored")
                run_notes.note_reason(run_notes.EMPTY, saved_name)
                any_failed = True
                continue
            size = os.path.getsize(path)
            extracted = extract_attachment_text(folder, saved_name)
            if not os.path.exists(path):
                any_failed = True  # dropped (password protected): not an attachment
                continue
            attachments.append({
                "file_name": saved_name,
                "content_type": None,
                "size_bytes": size,
            })
            any_failed = any_failed or not extracted

    return attachments, any_failed

def listing_page_failed(log, page: int, exc: BaseException, found: int, problems: list | None) -> bool:
    """
    For a loop that walks a tender listing page by page. Call it from the
    `except` around one page's request. Returns True if the walk should stop
    here and carry on with the `found` tenders from the earlier pages: that is
    a later page (never the first) failing in a way that usually passes by
    itself -- a timeout, a dropped connection or a 500/502/503/504 that
    survived every retry. The page is noted in `problems` so the caller can
    report the listing as not read in full.

    Returns False for everything else (the first page, a 403/429, a block page,
    a bug), and the caller re-raises as before: those still stop the run.
    """
    if page <= 1 or not is_transient_error(exc):
        return False
    reason = (str(exc).strip().splitlines() or [""])[0][:200]
    log.error("listing: page %d could not be read (%s: %s) -- carrying on with the %d tender(s) "
              "found on the pages before it", page, type(exc).__name__, reason, found)
    if problems is not None:
        problems.append(f"page {page}")
    return True


def note_listing_problems(log, problems: list) -> bool:
    """Say so on the dashboard when a listing was not read in full. True if it was not."""
    if not problems:
        return False
    log.error("listing: not read in full (%s) -- the tenders on those pages and after them "
              "were not seen", ", ".join(problems))
    run_notes.note_late("listing", run_notes.LISTING, ", ".join(problems))
    return True


# ---------------------------------------------------------------------------
# A site that is down: start again once, and stop when it dies part-way
# ---------------------------------------------------------------------------

# Seconds to wait before each extra try of a whole run that found the site
# unreachable (SITE_TOTAL_FAILURE with nothing scraped). Two extra tries, about
# four minutes in all: long enough for a restart or a network blip to pass,
# short enough that a site that is really down costs little. () turns it off.
RUN_RETRY_WAITS = (60, 180)


# How many login forms have been posted in this process. The whole-run retry
# below reads it: a run that already sent the credentials is not run again,
# because that would post them again (a lockout risk) -- the site was reachable,
# whatever failed after the login was not "unreachable".
_LOGIN_ATTEMPTS = 0


def note_login_attempt() -> None:
    """Call just before posting a login form (submit_login does it itself)."""
    global _LOGIN_ATTEMPTS
    _LOGIN_ATTEMPTS += 1


def login_attempts() -> int:
    return _LOGIN_ATTEMPTS


def retry_if_site_unreachable(run_once, log, waits=None):
    """
    run_once() -> (code, tenders). Runs it, and runs it again after each wait in
    `waits` while it ends in SITE_TOTAL_FAILURE with no tenders scraped: the site
    could not be reached at all (the first page, the login page), so nothing was
    kept and nothing is repeated. Any other result is returned as it is, and so is
    the last one when every try fails.

    For the scrapers without a browser to restart (GrantConnect, AusTender, NT,
    buy.nsw, eTender, WA, TAS). A failed login, a block, a rate limit and a
    structure change are never retried -- and nor is a run that already posted
    its login form (see note_login_attempt): the credentials are sent once.
    """
    waits = RUN_RETRY_WAITS if waits is None else waits
    logins_before = login_attempts()
    code, tenders = run_once()
    for attempt, wait in enumerate(waits, start=2):
        if code != SITE_TOTAL_FAILURE or tenders:
            break
        if login_attempts() > logins_before:
            log.warning("the run failed after the login was sent -- not trying the whole "
                        "run again, so the credentials are not posted a second time")
            break
        log.warning("the site could not be reached -- trying the whole run again in %ds "
                    "(try %d of %d)", wait, attempt, len(waits) + 1)
        time.sleep(wait)
        code, tenders = run_once()
    return code, tenders


def level_unparsed_pages(log, unparsed: list, attempted: int, tender_codes: set) -> None:
    """
    For a loop that reports a detail page it could not parse as SITE_STRUCTURE_CHANGE.
    `unparsed` holds those tenders' refs, `attempted` how many tenders the loop
    tried in all. When EVERY page failed to parse the layout has changed and the
    site is reported so. When only some did (a tender withdrawn between the
    listing and its page, a page that did not render in time) those tenders are
    reported as incomplete -- fetched again next run -- rather than failing the
    whole site over one page.
    """
    if not unparsed:
        return
    if len(unparsed) >= attempted:
        tender_codes.add(SITE_STRUCTURE_CHANGE)
        return
    log.warning("%d of %d tender page(s) could not be parsed -- reported as incomplete "
                "(nothing stored for them, so they are tried again next run), not as a "
                "site-wide structure change, since the other pages parsed",
                len(unparsed), attempted)
    for ref in unparsed:
        run_notes.note_late(ref, run_notes.PAGE)
    tender_codes.add(TENDER_PARTIAL)


# A site that answers nothing for this many tenders in a row has stopped
# answering. Each tender spends its retries first, so carrying on only burns the
# run's time (hours, over a few hundred tenders) for nothing.
DEAD_SITE_AFTER = 15


class DeadSiteWatch:
    """
    Count tenders in a row that failed with a transient error (timeout, dropped
    connection, 500/502/503/504 that survived its retries).

        watch = common.DeadSiteWatch()
        for ...:
            try:
                ...
                watch.succeeded()
            except Exception as exc:
                if watch.failed(exc):       # True once the site looks dead
                    break

    Any tender that gets an answer, or fails for another reason, starts the count
    again. `stopped` says the loop ended because of this.
    """

    def __init__(self, limit: int | None = None):
        self.limit = DEAD_SITE_AFTER if limit is None else limit
        self.streak = 0
        self.stopped = False

    def succeeded(self) -> None:
        self.streak = 0

    def failed(self, exc: BaseException) -> bool:
        self.streak = self.streak + 1 if is_transient_error(exc) else 0
        if self.limit and self.streak >= self.limit:
            self.stopped = True
        return self.stopped

    def report(self, log, done: int, total: int) -> None:
        """Log it and put it on the dashboard. Call once, after the loop."""
        log.error("the site stopped answering: the last %d tender(s) all timed out or "
                  "failed with a server error -- stopping after %d of %d; the rest are "
                  "left for the next run", self.streak, done, total)
        run_notes.note_late(
            "site", run_notes.SITE_STOPPED,
            f"stopped after {done} of {total} tenders")


class BrowserLostError(Exception):
    """Raised when the browser itself has died (crashed, closed, lost its session)."""


# Fragments of the messages Selenium/Chrome give when the browser is gone, as
# opposed to one page misbehaving. Matched on text so common.py needs no selenium import.
_BROWSER_GONE_MARKERS = (
    "invalid session id",
    "session deleted",
    "chrome not reachable",
    "disconnected",
    "no such window",
    "target window already closed",
    "web view not found",
    "not connected to devtools",
    "browser has closed",
    "connection refused",
    "max retries exceeded",
)


def browser_is_gone(exc: BaseException) -> bool:
    """True if `exc` means the browser died, so every later page would fail the same way."""
    text = f"{type(exc).__name__}: {exc}".lower()
    return any(marker in text for marker in _BROWSER_GONE_MARKERS)


#The site-level code for a run, from the per-tender codes it collected. Same
#precedence every scraper uses: a structure change outranks a partial tender,
#since it means the scraper itself needs fixing.
def site_code_from(tender_codes) -> int:
    codes = set(tender_codes) - {SITE_SUCCESS}
    if not codes:
        return SITE_SUCCESS
    for code in (SITE_BOT_BLOCKED, SITE_STRUCTURE_CHANGE, SITE_LOGIN_FAILED, SITE_RATE_LIMITED):
        if code in codes:
            return code
    return TENDER_PARTIAL

#-----
#This is our helper function to log in to sites
#-----

def submit_login(client: httpx.Client, url: str, payload: dict) -> httpx.Response:
    note_login_attempt()
    return client.post(
        url,
        data=payload,
        headers={
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/120.0.0.0 Safari/537.36"
            ),
        },
        follow_redirects=True,
        timeout=25.0,
    )