"""
Two things the 5am WA run showed, both in error_scrapers/common.py:

* Junk that rides along in a document package (an Office "~$" lock file, an empty
  TRIMfiles.dat) made a whole tender TENDER_PARTIAL. They are skipped now.
* Four .xlsX files in DHW0357125 failed in openpyxl over a cell style with no name,
  and the tender was partial. The cell values are read directly now.
"""

import io
import re
import zipfile

import openpyxl
import pytest

from error_scrapers import common, run_notes


def _unpack(tmp_path, members: dict):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        for name, data in members.items():
            zf.writestr(name, data)
    buffer.seek(0)
    return common.unpack_zip(buffer, str(tmp_path))


def _nameless_style_workbook(path):
    """A real workbook whose cell style has lost its name, as the portal's files had."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bill 5"
    ws.append(["Item", "Description", "Qty", "Rate"])
    ws.append([1, "Provisional sum", 2, 150.5])
    wb.create_sheet("Notes")["B2"] = "Allowance"
    good = io.BytesIO()
    wb.save(good)
    good.seek(0)
    with zipfile.ZipFile(good) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/styles.xml":
                data = re.sub(rb'(<cellStyle )name="[^"]*" ', rb"\1", data)
            zout.writestr(item, data)


# --- junk in a package -------------------------------------------------------

def test_an_office_lock_file_is_ignored_not_a_failure(tmp_path):
    run_notes.reset()
    attachments, any_failed = _unpack(tmp_path, {
        "~$618-ST-RPT-01_A Column Inspection Report.pdf": b"\x05Ben Gilmore" + b" " * 150,
        "Report.txt": b"hello"})
    assert any_failed is False
    assert [a["file_name"] for a in attachments] == ["Report.txt"]
    assert not (tmp_path / "~$618-ST-RPT-01_A Column Inspection Report.pdf").exists()
    assert run_notes._pending == []  # nothing is noted against the tender


def test_macos_and_windows_leftovers_are_ignored(tmp_path):
    attachments, any_failed = _unpack(tmp_path, {
        "__MACOSX/._Report.txt": b"resource fork", ".DS_Store": b"x", "Thumbs.db": b"x",
        "desktop.ini": b"x", "Report.txt": b"hello"})
    assert any_failed is False
    assert [a["file_name"] for a in attachments] == ["Report.txt"]


def test_an_empty_system_file_is_ignored_not_a_failure(tmp_path):
    attachments, any_failed = _unpack(tmp_path, {"TRIMfiles.dat": b"", "Report.txt": b"hello"})
    assert any_failed is False
    assert [a["file_name"] for a in attachments] == ["Report.txt"]
    assert not (tmp_path / "TRIMfiles.dat").exists()


@pytest.mark.parametrize("name", ["empty.pdf", "empty.docx", "empty.xlsx"])
def test_an_empty_document_is_still_a_download_problem(tmp_path, name):
    attachments, any_failed = _unpack(tmp_path, {name: b"", "Report.txt": b"hello"})
    assert any_failed is True
    assert [a["file_name"] for a in attachments] == ["Report.txt"]


def test_a_real_document_with_a_tilde_or_dollar_in_its_name_is_kept(tmp_path):
    attachments, any_failed = _unpack(tmp_path, {"Budget ~$5k.txt": b"hello", "Cost$.txt": b"hi"})
    assert any_failed is False
    assert sorted(a["file_name"] for a in attachments) == ["Budget ~$5k.txt", "Cost$.txt"]


def test_is_junk_file():
    assert common.is_junk_file("~$report.docx")
    assert common.is_junk_file("sub/folder/~$report.docx")
    assert common.is_junk_file("__MACOSX/anything.pdf")
    assert common.is_junk_file("sub\\.DS_Store")
    assert not common.is_junk_file("Report ~$5.pdf")
    assert not common.is_junk_file("Report.pdf")


# --- a spreadsheet openpyxl cannot open --------------------------------------

def test_openpyxl_really_does_refuse_the_nameless_style_workbook(tmp_path):
    path = tmp_path / "Bill.xlsX"
    _nameless_style_workbook(path)
    with pytest.raises(common.ExtractionError):
        common.extract_xlsx(str(path))


def test_a_workbook_openpyxl_refuses_is_read_cell_by_cell(tmp_path, caplog):
    path = tmp_path / "Bill.xlsX"
    _nameless_style_workbook(path)
    with caplog.at_level("WARNING", logger="scraper.common"):
        text = common.extract_with_fallback(common.extract_xlsx, str(path))
    assert "## Sheet: Bill 5" in text and "Provisional sum" in text and "150.5" in text
    assert "## Sheet: Notes" in text and "\tAllowance" in text  # B2 keeps its column
    assert "kept its cell values only" in caplog.text


def test_such_a_workbook_in_a_package_is_stored_with_its_text_and_not_partial(tmp_path):
    source = tmp_path / "src.xlsX"
    _nameless_style_workbook(source)
    out = tmp_path / "out"
    out.mkdir()
    attachments, any_failed = _unpack(out, {"Glen Forrest PS - Bill No. 05.xlsX": source.read_bytes()})
    assert any_failed is False
    assert [a["file_name"] for a in attachments] == ["Glen Forrest PS - Bill No. 05.xlsX"]
    assert "Provisional sum" in (out / "Glen Forrest PS - Bill No. 05.xlsX.txt").read_text(encoding="utf-8")


def test_the_cell_reader_matches_openpyxl_on_a_normal_workbook(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws.append(["Item", None, "Rate"])
    ws.append([1, "Concrete", 99.5])
    ws.append([None, None, None])
    ws.append([True, "Total", 100])
    wb.create_sheet("Empty")
    path = tmp_path / "ok.xlsx"
    wb.save(path)
    assert common.extract_xlsx_cells(str(path)) == common.extract_xlsx(str(path))


def test_the_cell_reader_handles_rich_text_inline_strings_and_gaps(tmp_path):
    path = tmp_path / "hand.xlsx"
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("xl/workbook.xml",
                    f'<workbook xmlns="{main}" xmlns:r="{rel}"><sheets>'
                    f'<sheet name="Hand" sheetId="1" r:id="rId1"/></sheets></workbook>')
        zf.writestr("xl/_rels/workbook.xml.rels",
                    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                    '<Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>')
        zf.writestr("xl/sharedStrings.xml",
                    f'<sst xmlns="{main}"><si><r><t>Rich </t></r><r><t>text</t></r><rPh><t>x</t></rPh></si></sst>')
        zf.writestr("xl/worksheets/sheet1.xml",
                    f'<worksheet xmlns="{main}"><sheetData><row r="1">'
                    '<c r="A1" t="s"><v>0</v></c>'
                    '<c r="C1" t="inlineStr"><is><t>inline</t></is></c>'
                    '<c r="D1" t="b"><v>1</v></c></row></sheetData></worksheet>')
    assert common.extract_xlsx_cells(str(path)) == "## Sheet: Hand\nRich text\t\tinline\tTrue"


def test_a_file_that_is_not_a_workbook_still_fails(tmp_path):
    path = tmp_path / "garbage.xlsx"
    path.write_bytes(b"PK\x03\x04 truncated")
    with pytest.raises(common.ExtractionError):
        common.extract_with_fallback(common.extract_xlsx, str(path))
