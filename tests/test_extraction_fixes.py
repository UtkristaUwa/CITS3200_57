"""
Tests for the fixes made after the 8 Oct 2026 production run:
  * DOCX tables with irregular merged cells / many rows no longer crash or crawl
  * GrantConnect and buy.nsw extract through the shared time limit
  * the NT login page fetch is retried on a timeout, the sign-in POST is not
"""

import time

import docx
import httpx
import pytest
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from error_scrapers import common


# --- DOCX tables -----------------------------------------------------------

def _set_text(table, row_index, cell_index, text):
    """Set a cell's text through its own XML, never through row.cells / table.cell,
    so building a test table is fast even on python-docx 1.1.0 (the production pin)."""
    from docx.table import _Cell
    _Cell(table._tbl.tr_lst[row_index].tc_lst[cell_index], table).text = text


def _save(doc, tmp_path, name="t.docx"):
    path = tmp_path / name
    doc.save(str(path))
    return str(path)


def test_docx_table_with_merged_cells_extracts(tmp_path):
    doc = docx.Document()
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2))      # horizontal merge across the header row
    table.cell(1, 0).merge(table.cell(2, 0))      # vertical merge
    _set_text(table, 0, 0, "Header spans all")
    _set_text(table, 1, 0, "Left tall")
    _set_text(table, 1, 1, "mid")
    _set_text(table, 2, 1, "end")
    text = common.extract_docx(_save(doc, tmp_path))
    assert "Header spans all" in text
    assert "Left tall" in text and "mid" in text and "end" in text


def test_docx_table_with_a_ragged_row_does_not_raise(tmp_path):
    """A row with fewer cells than the grid made python-docx's row.cells raise IndexError."""
    doc = docx.Document()
    table = doc.add_table(rows=2, cols=3)
    _set_text(table, 0, 0, "a")
    _set_text(table, 1, 0, "b")
    # drop the last cell of row 2 so it is shorter than the table grid
    tr = table._tbl.tr_lst[1]
    tr.remove(tr.tc_lst[-1])
    text = common.extract_docx(_save(doc, tmp_path))
    assert "a" in text and "b" in text


def test_docx_table_with_a_huge_grid_span_is_fast(tmp_path):
    """
    A cell declaring a huge gridSpan made python-docx's row.cells build a
    million-entry grid per row (24s for 20 rows; real documents have far more
    rows), which is the kind of input that hung extraction. Reading each row's
    own cells never expands the span.
    """
    doc = docx.Document()
    table = doc.add_table(rows=20, cols=3)
    for i, tr in enumerate(table._tbl.tr_lst):
        span = OxmlElement("w:gridSpan")
        span.set(qn("w:val"), "1000000")
        tr.tc_lst[0].get_or_add_tcPr().append(span)
        _set_text(table, i, 1, f"row {i}")
    path = _save(doc, tmp_path)
    started = time.monotonic()
    text = common.extract_docx(path)
    assert time.monotonic() - started < 5
    assert "row 19" in text


# --- scrapers go through the extraction time limit -------------------------

def test_grantconnect_extraction_uses_the_time_limit(monkeypatch, tmp_path):
    from error_scrapers.grant_connect import scraper

    calls = []

    def fake_run_extractor(extractor, path, timeout=None):
        calls.append(path)
        raise common.ExtractionError("timed out after 300s")

    monkeypatch.setattr(common, "run_extractor", fake_run_extractor)

    class Resp:
        headers = {"content-type": "application/pdf"}
        def raise_for_status(self): pass
        def iter_bytes(self, *a, **k): yield b"%PDF-1.4"
        def iter_raw(self, *a, **k): yield b"%PDF-1.4"
        def __enter__(self): return self
        def __exit__(self, *a): return False

    class Client:
        def stream(self, *a, **k): return Resp()

    monkeypatch.setattr(
        common, "save_attachment_stream",
        lambda out, name, resp: _write(tmp_path, name))
    code, attachments = scraper.process_documents(
        Client(), [{"url": "http://x/a.pdf", "file_name": "a.pdf"}], str(tmp_path)
    ) if hasattr(scraper, "process_documents") else (None, None)
    if code is None:
        pytest.skip("process_documents has a different name here")
    assert calls, "run_extractor was not used"
    assert code == common.TENDER_PARTIAL


def _write(folder, name):
    path = folder / name
    path.write_bytes(b"%PDF-1.4")
    return str(path)


# --- NT login fetch --------------------------------------------------------

def test_nt_login_page_fetch_is_retried_but_the_post_is_not(monkeypatch):
    from error_scrapers.nt_qtol import scraper

    monkeypatch.setattr(common.time, "sleep", lambda s: None)
    monkeypatch.setattr(scraper, "credentials", lambda: ("user", "pass"))

    login_html = (
        "<form action='/Account/LogOn'><input name='UserId'/>"
        "<input type='hidden' name='tok' value='1'/></form>"
    )

    class Client:
        def __init__(self):
            self.gets = 0
            self.posts = 0

        def get(self, url, **kwargs):
            self.gets += 1
            if self.gets < 3:
                raise httpx.ReadTimeout("slow")
            return httpx.Response(200, text=login_html, request=httpx.Request("GET", url))

        def post(self, url, **kwargs):
            self.posts += 1
            raise httpx.ReadTimeout("slow")

    monkeypatch.setattr(scraper, "LOGIN_FIELD_SELECTOR", "input[name='UserId']")
    client = Client()
    with pytest.raises(httpx.TimeoutException):
        scraper.login(client)
    assert client.gets == 3       # the page fetch was retried
    assert client.posts == 1      # the sign-in was attempted once only


# --- the Word-file fixes found on python-docx 1.1.0 (what production runs) ---

def test_an_orphan_vertical_merge_cell_no_longer_raises(tmp_path):
    """
    A vertically merged cell with nothing above it made python-docx 1.1.0's
    row.cells raise "list index out of range" (the failure seen on three VIC
    contract files). Reading each row's own cells has no such case.
    """
    doc = docx.Document()
    table = doc.add_table(rows=3, cols=3)
    for i in range(3):
        _set_text(table, i, 1, f"keep {i}")
    for i in (1, 2):
        table._tbl.tr_lst[i].tc_lst[0].get_or_add_tcPr().append(OxmlElement("w:vMerge"))
    first = table._tbl.tr_lst[0]
    first.remove(first.tc_lst[0])
    text = common.extract_docx(_save(doc, tmp_path))
    assert "keep 2" in text


def test_a_few_hundred_row_table_is_fast(tmp_path):
    """python-docx 1.1.0 took 18s for 400 rows and 74s for 800 (quadratic)."""
    doc = docx.Document()
    table = doc.add_table(rows=400, cols=10)
    _set_text(table, 399, 9, "last cell")
    path = _save(doc, tmp_path)
    started = time.monotonic()
    text = common.extract_docx(path)
    assert time.monotonic() - started < 3
    assert "last cell" in text


def _docx_with_dangling_relationship(tmp_path):
    """A Word file python-docx can't open: a relationship points at a part that isn't there."""
    import zipfile
    doc = docx.Document()
    doc.add_paragraph("The surviving body text")
    good = _save(doc, tmp_path, "good.docx")
    bad = str(tmp_path / "bad.docx")
    with zipfile.ZipFile(good) as src, zipfile.ZipFile(bad, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "word/_rels/document.xml.rels":
                data = data.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rId99" '
                    b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
                    b'Target="NULL"/></Relationships>')
            dst.writestr(item, data)
    return bad


def test_a_docx_python_docx_cannot_open_falls_back_to_its_body_text(tmp_path):
    bad = _docx_with_dangling_relationship(tmp_path)
    with pytest.raises(common.ExtractionError):
        common.extract_docx(bad)                      # the production failure
    text = common.extract_with_fallback(common.extract_docx, bad)
    assert "The surviving body text" in text


def test_a_file_that_is_not_a_zip_still_raises(tmp_path):
    path = tmp_path / "fake.docx"
    path.write_bytes(b"<html>not a word file</html>")
    with pytest.raises(common.ExtractionError):
        common.extract_with_fallback(common.extract_docx, str(path))


def test_the_fallback_is_only_used_for_word_files(tmp_path):
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"not a pdf")
    with pytest.raises(common.ExtractionError):
        common.extract_with_fallback(common.extract_pdf, str(path))


# --- empty files in a package -----------------------------------------------

def test_an_empty_file_in_a_package_is_a_download_problem(tmp_path):
    import io
    import zipfile
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        zf.writestr("empty.pdf", b"")
        zf.writestr("note.txt", b"hello")
    buffer.seek(0)
    attachments, any_failed = common.unpack_zip(buffer, str(tmp_path))
    assert any_failed is True
    assert [a["file_name"] for a in attachments] == ["note.txt"]
    assert not (tmp_path / "empty.pdf").exists()
