"""One stuck or broken attachment must not stall the run, and spreadsheets read cleanly."""

import logging
import os
import re
import time
import zipfile

import openpyxl
import pytest

from error_scrapers import common


def _fine(path):
    return "some text"


def _raises(path):
    raise ValueError("bad file")


def _hangs(path):
    time.sleep(60)


def _dies(path):
    os._exit(1)


def test_a_normal_extraction_returns_its_text(tmp_path):
    assert common.run_extractor(_fine, str(tmp_path / "a.xlsx"), timeout=5) == "some text"


def test_an_extractor_error_becomes_an_extraction_error(tmp_path):
    with pytest.raises(common.ExtractionError, match="ValueError: bad file"):
        common.run_extractor(_raises, str(tmp_path / "a.xlsx"), timeout=5)


def test_a_stuck_extractor_is_stopped_after_the_time_limit(tmp_path):
    started = time.time()
    with pytest.raises(common.ExtractionError, match="timed out after 1s on a.xlsx"):
        common.run_extractor(_hangs, str(tmp_path / "a.xlsx"), timeout=1)
    assert time.time() - started < 10  # killed, not waited out


def test_an_extractor_that_dies_is_reported(tmp_path):
    with pytest.raises(common.ExtractionError, match="died without a result"):
        common.run_extractor(_dies, str(tmp_path / "a.xlsx"), timeout=5)


def test_a_large_result_comes_back_whole(tmp_path):
    big = "x" * 5_000_000
    assert common.run_extractor(lambda p: big, str(tmp_path / "a.xlsx"), timeout=30) == big


def test_a_stuck_file_is_a_failure_and_the_next_file_still_works(tmp_path, monkeypatch, caplog):
    monkeypatch.setattr(common, "EXTRACTION_TIMEOUT_SECONDS", 1)
    monkeypatch.setitem(common.EXTRACTORS, ".xlsx", _hangs)
    (tmp_path / "stuck.xlsx").write_bytes(b"x")
    (tmp_path / "ok.txt").write_text("hello", encoding="utf-8")
    with caplog.at_level(logging.INFO, logger="scraper.common"):
        assert common.extract_attachment_text(str(tmp_path), "stuck.xlsx") is False
        assert common.extract_attachment_text(str(tmp_path), "ok.txt") is True
    assert not os.path.exists(tmp_path / "stuck.xlsx.txt")
    assert os.path.exists(tmp_path / "ok.txt.txt")
    messages = [r.getMessage() for r in caplog.records]
    assert "extracting text from stuck.xlsx" in messages  # a hang can be traced to its file
    assert any("text extraction failed for stuck.xlsx" in m for m in messages)


def _huge_declared_sheet(path):
    """A workbook whose sheet claims to be A1:XFD1048576 but holds two rows."""
    small = str(path) + ".small"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["A2"], ws["B2"] = "Item", "Rate", "Design", 150
    wb.save(small)
    with zipfile.ZipFile(small) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]*"', b'<dimension ref="A1:XFD1048576"', data)
            zout.writestr(info, data)


def test_a_sheet_that_declares_a_huge_size_is_read_without_padding(tmp_path):
    path = tmp_path / "huge.xlsx"
    _huge_declared_sheet(path)
    text = common.extract_xlsx(str(path))
    assert text == "## Sheet: Sheet\nItem\tRate\nDesign\t150"
